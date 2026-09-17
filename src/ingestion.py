from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import fitz  # PyMuPDF
import pandas as pd
import pdfplumber
import pytesseract
from docx import Document as DocxDocument
from PIL import Image

from .config import settings

logger = logging.getLogger(__name__)

UnitType = Literal["text", "table", "image_caption", "ocr_text"]


@dataclass
class RawUnit:
    doc_id: str
    file_name: str
    page: int
    unit_type: UnitType
    content: str                 
    extra: dict = field(default_factory=dict)  


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def _page_has_text(page: "pdfplumber.page.Page", min_chars: int = 20) -> bool:
    text = page.extract_text() or ""
    return len(text.strip()) >= min_chars


def _ocr_page(pdf_path: str, page_number: int) -> str:
    """Rasterize a PDF page and run Tesseract OCR — used for scanned pages."""
    doc = fitz.open(pdf_path)
    page = doc.load_page(page_number)
    zoom = settings.ocr_dpi / 72
    mat = fitz.Matrix(zoom, zoom)
    pix = page.get_pixmap(matrix=mat)
    img = Image.open(io.BytesIO(pix.tobytes("png")))
    text = pytesseract.image_to_string(img, lang=settings.ocr_lang)
    doc.close()
    return text


def _extract_pdf_tables(page: "pdfplumber.page.Page") -> list[pd.DataFrame]:
    tables = []
    for raw_table in page.extract_tables():
        if raw_table and len(raw_table) > 1:
            df = pd.DataFrame(raw_table[1:], columns=raw_table[0])
            tables.append(df)
    return tables


def _extract_pdf_images(pdf_path: str, page_number: int) -> list[bytes]:
    doc = fitz.open(pdf_path)
    page = doc.load_page(page_number)
    images = []
    for img in page.get_images(full=True):
        xref = img[0]
        base = doc.extract_image(xref)
        images.append(base["image"])
    doc.close()
    return images


def ingest_pdf(path: str, doc_id: str, describe_images_fn=None) -> list[RawUnit]:
    
    file_name = Path(path).name
    units: list[RawUnit] = []

    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            page_num = i + 1

            # 1. Native text vs scanned page
            if _page_has_text(page):
                text = page.extract_text() or ""
                if text.strip():
                    units.append(RawUnit(doc_id, file_name, page_num, "text", text))
            else:
                ocr_text = _ocr_page(path, i)
                if ocr_text.strip():
                    units.append(
                        RawUnit(doc_id, file_name, page_num, "ocr_text", ocr_text)
                    )

            # 2. Tables (charts-as-tables, financial tables, invoice line items)
            for t_idx, df in enumerate(_extract_pdf_tables(page)):
                md = df.to_markdown(index=False)
                units.append(
                    RawUnit(
                        doc_id, file_name, page_num, "table",
                        f"Table {t_idx + 1} on page {page_num}:\n{md}",
                        extra={"table_df": df, "table_index": t_idx},
                    )
                )

            # 3. Embedded images / charts / receipts / invoices as images
            for img_idx, img_bytes in enumerate(_extract_pdf_images(path, i)):
                if describe_images_fn:
                    caption = describe_images_fn(img_bytes)
                else:
                    caption = (
                        f"[Image {img_idx + 1} on page {page_num} — "
                        f"no vision captioning configured]"
                    )
                units.append(
                    RawUnit(
                        doc_id, file_name, page_num, "image_caption", caption,
                        extra={"image_index": img_idx},
                    )
                )

    logger.info("Ingested %s: %d units", file_name, len(units))
    return units


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def ingest_docx(path: str, doc_id: str, describe_images_fn=None) -> list[RawUnit]:
    file_name = Path(path).name
    units: list[RawUnit] = []
    doc = DocxDocument(path)

    
    buf = []
    block_idx = 1
    for para in doc.paragraphs:
        if para.text.strip():
            buf.append(para.text)
        if len("\n".join(buf)) > 1500:  
            units.append(RawUnit(doc_id, file_name, block_idx, "text", "\n".join(buf)))
            buf, block_idx = [], block_idx + 1
    if buf:
        units.append(RawUnit(doc_id, file_name, block_idx, "text", "\n".join(buf)))

    # Tables
    for t_idx, table in enumerate(doc.tables):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        if len(rows) > 1:
            df = pd.DataFrame(rows[1:], columns=rows[0])
            md = df.to_markdown(index=False)
            units.append(
                RawUnit(
                    doc_id, file_name, 0, "table",
                    f"Table {t_idx + 1}:\n{md}",
                    extra={"table_df": df, "table_index": t_idx},
                )
            )

    # Images embedded in the docx package
    img_idx = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            img_bytes = rel.target_part.blob
            caption = (
                describe_images_fn(img_bytes)
                if describe_images_fn
                else f"[Embedded image {img_idx + 1} — no vision captioning configured]"
            )
            units.append(
                RawUnit(doc_id, file_name, 0, "image_caption", caption,
                         extra={"image_index": img_idx})
            )
            img_idx += 1

    logger.info("Ingested %s: %d units", file_name, len(units))
    return units


def ingest_file(path: str, doc_id: str, describe_images_fn=None) -> list[RawUnit]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return ingest_pdf(path, doc_id, describe_images_fn)
    elif ext == ".docx":
        return ingest_docx(path, doc_id, describe_images_fn)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
